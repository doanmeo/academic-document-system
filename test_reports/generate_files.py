import os
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import win32com.client as win32

def create_report_docx(filename, title, subject, student, advisor, problem_text, solution_text, tech_stack, results_text):
    doc = docx.Document()

    # Standard 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Header 1: Bộ & Trường
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(12)
    r1.bold = True

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("TRƯỜNG ĐẠI HỌC GIAO THÔNG VẬN TẢI")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(13)
    r2.bold = True

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run("KHOA CÔNG NGHỆ THÔNG TIN")
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(13)
    r3.font.color.rgb = RGBColor(0, 51, 102)
    r3.bold = True

    p_divider = doc.add_paragraph()
    p_divider.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_div = p_divider.add_run("---------------***---------------")
    r_div.font.name = "Times New Roman"
    r_div.font.size = Pt(10)

    doc.add_paragraph() # Empty line

    # Title
    p_badge = doc.add_paragraph()
    p_badge.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rb = p_badge.add_run("BÁO CÁO HỌC PHẦN / ĐỒ ÁN MÔN HỌC")
    rb.font.name = "Times New Roman"
    rb.font.size = Pt(14)
    rb.bold = True

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rt = p_title.add_run(f"ĐỀ TÀI: {title.upper()}")
    rt.font.name = "Times New Roman"
    rt.font.size = Pt(16)
    rt.font.color.rgb = RGBColor(0, 102, 204)
    rt.bold = True

    doc.add_paragraph() # Empty line

    # Metadata lines
    meta_items = [
        ("Học phần / Môn học:", subject),
        ("Sinh viên thực hiện:", student),
        ("Mã số sinh viên (MSV):", "231230756"),
        ("Chuyên ngành:", "Kỹ thuật Phần mềm (Khoa CNTT - UTC)"),
        ("Giảng viên hướng dẫn:", advisor),
        ("Năm học:", "2026 - 2027"),
    ]

    for label, val in meta_items:
        p_m = doc.add_paragraph()
        p_m.paragraph_format.space_after = Pt(2)
        r_l = p_m.add_run(label + " ")
        r_l.font.name = "Times New Roman"
        r_l.font.size = Pt(12)
        r_l.bold = True
        
        r_v = p_m.add_run(val)
        r_v.font.name = "Times New Roman"
        r_v.font.size = Pt(12)

    doc.add_paragraph() # Empty line

    # Section I
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(4)
    rh1 = h1.add_run("I. ĐẶT VẤN ĐỀ VÀ TÍNH CẤP THIẾT CỦA ĐỀ TÀI")
    rh1.font.name = "Times New Roman"
    rh1.font.size = Pt(13)
    rh1.font.color.rgb = RGBColor(0, 51, 102)
    rh1.bold = True

    p_prob = doc.add_paragraph()
    p_prob.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rp = p_prob.add_run(problem_text)
    rp.font.name = "Times New Roman"
    rp.font.size = Pt(12)

    # Section II
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)
    rh2 = h2.add_run("II. GIẢI PHÁP KỸ THUẬT VÀ CÔNG NGHỆ ÁP DỤNG")
    rh2.font.name = "Times New Roman"
    rh2.font.size = Pt(13)
    rh2.font.color.rgb = RGBColor(0, 51, 102)
    rh2.bold = True

    p_sol = doc.add_paragraph()
    p_sol.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rs = p_sol.add_run(solution_text)
    rs.font.name = "Times New Roman"
    rs.font.size = Pt(12)

    p_tech = doc.add_paragraph()
    p_tech.paragraph_format.space_before = Pt(4)
    rt_l = p_tech.add_run("• Công nghệ chủ chốt: ")
    rt_l.font.name = "Times New Roman"
    rt_l.font.size = Pt(12)
    rt_l.bold = True
    rt_v = p_tech.add_run(tech_stack)
    rt_v.font.name = "Times New Roman"
    rt_v.font.size = Pt(12)

    # Section III
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(4)
    rh3 = h3.add_run("III. KẾT QUẢ THỰC NGHIỆM VÀ KẾT LUẬN")
    rh3.font.name = "Times New Roman"
    rh3.font.size = Pt(13)
    rh3.font.color.rgb = RGBColor(0, 51, 102)
    rh3.bold = True

    p_res = doc.add_paragraph()
    p_res.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rr = p_res.add_run(results_text)
    rr.font.name = "Times New Roman"
    rr.font.size = Pt(12)

    doc.save(filename)

base_dir = os.path.abspath("academic-document-system/test_reports")
os.makedirs(base_dir, exist_ok=True)

f1_docx = os.path.join(base_dir, "Bao_Cao_Giam_Sat_Ngap_Lut_IoT.docx")
f1_pdf = os.path.join(base_dir, "Bao_Cao_Giam_Sat_Ngap_Lut_IoT.pdf")

f2_docx = os.path.join(base_dir, "Bao_Cao_Tro_Ly_Ao_RAG_UTC.docx")
f2_pdf = os.path.join(base_dir, "Bao_Cao_Tro_Ly_Ao_RAG_UTC.pdf")

# Generate File 1
create_report_docx(
    f1_docx,
    title="Hệ thống Cảnh báo Ngập lụt Đô thị Thông minh qua Mạng Cảm biến IoT",
    subject="Mạng máy tính & Hệ thống Nhúng",
    student="Vũ Đức An Ninh",
    advisor="PGS. TS. Trần Văn Hùng",
    problem_text="Tình trạng ngập úng đô thị trong mùa mưa bão tại các thành phố lớn như Hà Nội đang diễn biến phức tạp, gây ùn tắc giao thông nghiêm trọng và thiệt hại lớn về tài sản của người dân. Các phương pháp quan trắc truyền thống phụ thuộc vào kiểm tra thủ công, chậm trễ trong việc cập nhật cảnh báo và chưa có giải pháp tự động gửi lộ trình an toàn cho người tham gia giao thông.",
    solution_text="Đề tài đề xuất giải pháp thiết kế mạng lưới trạm đo mực nước tự động ứng dụng cảm biến sóng siêu âm không tiếp xúc, kết hợp bộ điều khiển ESP32 và công nghệ truyền tin tầm xa LoRaWAN. Dữ liệu ngập lụt được gửi định kỳ 30 giây một lần về Cloud Server qua giao thức MQTT, tự động số hóa và trực quan hóa theo thời gian thực trên bản đồ GIS.",
    tech_stack="ESP32, Cảm biến Siêu âm US-100, LoRaWAN SX1278, MQTT Broker EMQX, PostgreSQL/PostGIS, Leaflet GIS",
    results_text="Thực nghiệm tại 5 điểm đen ngập lụt khu vực Cầu Giấy cho thấy hệ thống hoạt động ổn định với độ trễ dưới 1.2 giây, tỷ lệ truyền gói tin thành công đạt 98.4%. Hệ thống phát cảnh báo tự động qua ứng dụng Web và Telegram Bot khi mực nước vượt ngưỡng 20cm, giúp người dân chủ động thay đổi lộ trình an toàn."
)

# Generate File 2
create_report_docx(
    f2_docx,
    title="Xây dựng Trợ lý Học tập Ảo giải đáp Quy chế Đào tạo UTC ứng dụng Kỹ thuật RAG",
    subject="Trí tuệ Nhân tạo & Xử lý Ngôn ngữ Tự nhiên",
    student="Vũ Đức An Ninh",
    advisor="ThS. Đào Thị Lệ Thủy",
    problem_text="Quy chế đào tạo theo hệ thống tín chỉ của Trường Đại học Giao thông Vận tải (UTC) bao gồm nhiều văn bản quy phạm phức tạp với hàng trăm điều khoản liên quan đến điều kiện xét tốt nghiệp, đăng ký học vượt, cải thiện điểm và xử lý học vụ. Sinh viên các khóa mới thường gặp khó khăn trong việc tiếp cận thông tin chính xác, dẫn đến tình trạng vi phạm quy chế hoặc đăng ký môn học không tối ưu.",
    solution_text="Xây dựng hệ thống Chatbot trợ lý ảo ứng dụng kỹ thuật Retrieval-Augmented Generation (RAG). Toàn bộ sổ tay sinh viên và văn bản quy chế đào tạo UTC được số hóa, làm sạch và phân mảnh ngữ nghĩa (Semantic Chunking) để lưu vào Cơ sở dữ liệu Vector. Khi sinh viên đặt câu hỏi, mô hình tìm kiếm ngữ nghĩa sẽ truy xuất các đoạn văn bản pháp quy chính xác nhất và làm ngữ cảnh cho Mô hình Ngôn ngữ Lớn (LLM) sinh câu trả lời kèm dẫn chứng điều khoản cụ thể.",
    tech_stack="Python, LangChain, Qdrant Vector Database, mô hình nhúng Text-Embedding-3, Google Gemini / Llama-3, React 19",
    results_text="Hệ thống đã lập chỉ mục hơn 150 trang văn bản quy chế chính thức của Trường ĐH Giao thông Vận tải. Thử nghiệm trên tập 200 câu hỏi học vụ thực tế đạt độ chính xác trả lời 94.5%, thời gian phản hồi trung bình 1.1 giây, khắc phục triệt để hiện tượng 'ảo giác' (hallucination) và hỗ trợ sinh viên tra cứu quy chế 24/7."
)

# Use MS Word COM to officially resave both DOCX and export native PDFs
try:
    word = win32.Dispatch("Word.Application")
    word.Visible = False

    for docx_path, pdf_path in [(f1_docx, f1_pdf), (f2_docx, f2_pdf)]:
        doc = word.Documents.Open(docx_path)
        # Resave DOCX natively via Word engine
        doc.Save()
        # Export PDF natively via Word engine
        doc.ExportAsFixedFormat(pdf_path, 17) # 17 = wdExportFormatPDF
        doc.Close(False)

    word.Quit()
    print("ALL FILES RESAVED AND EXPORTED NATIVELY BY MICROSOFT WORD!")
except Exception as e:
    print("Word COM resave warning:", e)
